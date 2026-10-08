#!/usr/bin/env python3
"""Generate the public, synthetic Spanish architecture document."""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "arquitectura-copilot-azure-openai.docx"
NAVY = (28, 48, 76)
TEAL = (22, 105, 124)
PALE = (233, 243, 247)
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
REVIEW_DATE = "2026-10-08"


def diagram(boxes, edges, width=1400, height=550):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, 21)
    small = ImageFont.truetype(FONT, 17)
    for x1, y1, x2, y2, label in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=14, fill=PALE, outline=TEAL, width=3)
        lines = label.split("\n")
        for index, line in enumerate(lines):
            w = draw.textlength(line, font=font)
            draw.text(((x1 + x2 - w) / 2, (y1 + y2) / 2 - 14 * len(lines) + index * 29),
                      line, font=font, fill=NAVY)
    for x1, y1, x2, y2, label in edges:
        draw.line((x1, y1, x2, y2), fill=TEAL, width=4)
        if x2 > x1:
            tip = [(x2, y2), (x2 - 12, y2 - 7), (x2 - 12, y2 + 7)]
        elif x2 < x1:
            tip = [(x2, y2), (x2 + 12, y2 - 7), (x2 + 12, y2 + 7)]
        else:
            tip = [(x2, y2), (x2 - 7, y2 - 12), (x2 + 7, y2 - 12)]
        draw.polygon(tip, fill=TEAL)
        if label:
            w = draw.textlength(label, font=small)
            cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
            draw.rectangle((cx - w / 2 - 7, cy - 24, cx + w / 2 + 7, cy + 1), fill="white")
            draw.text((cx - w / 2, cy - 22), label, font=small, fill=NAVY)
    buffer = BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    return buffer


def paragraph(doc, text="", style=None):
    return doc.add_paragraph(text, style=style)


def table(doc, headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = "Light Shading Accent 1"
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(result.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(result.add_row().cells, row):
            cell.text = value
    doc.add_paragraph()


def picture(doc, image, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(image, width=Cm(16.4))
    paragraph(doc, caption, "Caption")


def heading(doc, title, level=1):
    doc.add_heading(title, level)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.1)
    sec.left_margin = sec.right_margin = Cm(2.2)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor(*NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.color.rgb = RGBColor(*TEAL)
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.italic = True
    props = doc.core_properties
    props.title = "Arquitectura y posición de seguridad - Copilot con Azure OpenAI"
    props.subject = "Arquitectura pública declarada, controles y brechas"
    props.author = "Proyecto Copilot con Azure OpenAI"
    props.last_modified_by = "Proyecto Copilot con Azure OpenAI"
    props.keywords = "arquitectura;seguridad;Azure OpenAI;Key Vault"
    props.comments = ""
    props.created = props.modified = datetime(2026, 10, 8, tzinfo=timezone.utc)

    doc.add_heading("Arquitectura del proyecto", 0)
    paragraph(doc, "Copilot CLI y VS Code con modelos propios en Azure OpenAI", "Subtitle")
    paragraph(doc, f"Edición pública | Revisión: {REVIEW_DATE} | Escenario principal: endpoints públicos")
    paragraph(doc, "Documento propio del proyecto. Describe código y plantillas, no una auditoría de "
              "suscripciones, equipos ni sesiones reales. Sin identificadores ni datos operativos.")

    heading(doc, "1. Alcance y criterio de evidencia")
    paragraph(doc, "El repositorio contiene un cliente Python para GitHub Copilot CLI en Linux/WSL, "
              "una integración opcional con VS Code y aprovisionamiento Azure por Bicep. Reutilizar "
              "cuenta y vault existentes es el valor por defecto. La infraestructura nueva, los "
              "deployments y el presupuesto son opcionales y requieren plan y aprobación.")
    table(doc, ("Etiqueta", "Qué acredita", "Qué no acredita"), [
        ("Implementado", "Comportamiento observable en código, scripts o plantilla.",
         "No prueba ejecución ni configuración efectiva en Azure."),
        ("Declarado", "Valor deseado para recursos creados con esta plantilla.",
         "No impone controles a recursos reutilizados."),
        ("Pendiente", "Requiere evidencia privada y prueba reproducible.",
         "No debe publicarse como cumplimiento."),
    ])
    paragraph(doc, "Las guías externas aportadas (Guia_acceso_publico_Copilot_Azure_OpenAI e "
              "Informe_seguridad_Copilot_Azure_OpenAI) se usaron como criterios de contraste, "
              "no como arquitectura propia ni como evidencia del despliegue. El informe histórico "
              "evaluaba otra revisión; sus hallazgos H01-H06 no se atribuyen automáticamente al cliente actual.")

    heading(doc, "2. Contexto, componentes y límites de confianza")
    picture(doc, diagram([
        (40, 90, 325, 190, "Persona + equipo\nLinux / WSL / VS Code"),
        (410, 40, 745, 140, "Copilot CLI + lanzador\nPython / JSON local"),
        (410, 255, 745, 355, "VS Code / importador\nsecreto nativo del editor"),
        (850, 40, 1170, 140, "Microsoft Entra + ARM\nidentidad / metadatos"),
        (850, 220, 1170, 320, "Azure Key Vault\nconfiguración y clave"),
        (850, 400, 1170, 500, "Azure OpenAI\ndeployments / inferencia"),
    ], [
        (325, 115, 410, 115, "uso"), (325, 160, 410, 305, ""),
        (745, 85, 850, 85, "Entra/ARM"), (745, 130, 850, 267, ""),
        (745, 320, 850, 275, "lectura"),
        (745, 125, 850, 440, ""), (745, 350, 850, 445, ""),
    ], height=550), "Figura 1. Componentes lógicos (flechas = dependencias; no prueba de tráfico efectivo).")
    paragraph(doc, "La frontera local contiene perfil no secreto, procesos, cachés de autenticación y "
              "datos del workspace. Entra y ARM son plano de identidad y control; Key Vault es plano "
              "de secretos; Azure OpenAI es plano de inferencia. El firewall de los recursos Azure "
              "no restringe las lecturas del workspace ni la salida hacia GitHub, extensiones, MCP "
              "u otras herramientas; estas rutas necesitan controles propios del puesto y tenant.")
    paragraph(doc, "VS Code usa su servicio de credenciales mediante grupo Custom Endpoint; el "
              "importador experimental depende de comando interno y versión del editor. No se "
              "copian bases de datos de credenciales ni se fabrican referencias. Los subagentes "
              "pueden heredar el modelo de la sesión; utilidades internas pueden pedir otro "
              "deployment, cuya ruta efectiva exige comprobación independiente.")

    heading(doc, "3. Secuencias críticas y manejo de secretos")
    picture(doc, diagram([
        (45, 65, 345, 155, "Operador\nmanifiesto privado"),
        (515, 65, 835, 155, "plan / what-if\naprobación vinculada"),
        (1020, 65, 1330, 155, "ARM / Bicep\nrecursos opcionales"),
        (45, 300, 345, 390, "Lanzador\nperfil JSON"),
        (515, 240, 835, 330, "Entra + ARM\nvalidación identidad"),
        (515, 385, 835, 475, "Key Vault\nconfig + clave"),
        (1020, 300, 1330, 390, "Copilot CLI\nAzure OpenAI"),
    ], [
        (345, 110, 515, 110, "plan"), (835, 110, 1020, 110, "apply"),
        (345, 340, 515, 285, "login"), (345, 375, 515, 430, ""),
        (515, 455, 345, 375, "clave"),
        (345, 335, 1020, 335, "clave via entorno"),
    ], height=540), "Figura 2. Secuencias resumidas: aprovisionamiento (arriba) y arranque (abajo).")
    table(doc, ("Paso", "Control existente en el repositorio", "Límite o verificación"), [
        ("Plan y apply", "Preflight, what-if, hash de manifiesto/plantilla, confirmación; sin claves en Bicep.",
         "Comprobar drift y recursos parcialmente creados; informes privados."),
        ("Bootstrap", "seed_vault.py escribe configuración y clave por HTTPS; permisos de escritura separados.",
         "No rota clave ni sustituye un secreto ya existente automáticamente."),
        ("Arranque", "foundry.py valida perfil JSON, identidad/endpoint contra ARM y TPM antes de leer clave.",
         "doctor no lee clave ni infiere; smoke-test directo no prueba ruta Copilot."),
        ("Inferencia", "Entra lee el vault; Copilot recibe API key en entorno, no en argumento ni archivo compartido.",
         "Clave puede permanecer en proceso; revocar RBAC del vault no revoca una clave copiada."),
        ("Editor", "Importación optativa, consentimiento y almacén nativo de secretos de VS Code.",
         "Verificar versión, persistencia y chat/Agent real con datos sintéticos."),
    ])

    heading(doc, "4. Decisiones arquitectónicas y de seguridad")
    table(doc, ("ID", "Decisión y motivo", "Consecuencia / riesgo residual"), [
        ("D01", "Endpoint público con IPs de salida aprobadas cuando son estables.",
         "No es red privada; clientes fuera de la lista fallan."),
        ("D02", "Perfil público abierto solo con aceptación explícita del riesgo.",
         "Una clave robada puede usarse desde cualquier IP hasta su revocación."),
        ("D03", "Entra + RBAC de lectura de Key Vault; clave API para inferencia.",
         "disableLocalAuth=false; no atribución individual Entra en inferencia."),
        ("D04", "Configurar metadatos en JSON validado y secreto en vault; no ejecutar .env.",
         "Protección de hijos mediante filtrado, no sandbox frente al mismo usuario/root."),
        ("D05", "Reutilizar recursos; crear solo por plan aprobado.",
         "Plantilla no impone reglas a recursos reutilizados."),
        ("D06", "Elegir SKU/modelo con catálogo y cuota revisados; consentimiento para GlobalStandard.",
         "La región del recurso no garantiza geografía de proceso en SKU global."),
        ("D07", "Mantener secretos del editor dentro del mecanismo nativo.",
         "La automatización de importación es experimental y dependiente de versión."),
    ])

    heading(doc, "5. Topología de red pública y alternativa privada")
    table(doc, ("Perfil", "Azure OpenAI y Key Vault", "Juicio"), [
        ("public_selected_ips", "publicNetworkAccess=Enabled; defaultAction=Deny; IPs IPv4 aprobadas en ambos.",
         "Preferido si hay salida estable; verificar configuración efectiva y IP de salida."),
        ("public_any_ip", "publicNetworkAccess=Enabled; defaultAction=Allow; autenticación obligatoria.",
         "Mayor exposición; aprobación de riesgo, respuesta y monitoreo."),
        ("Red privada (futuro)", "Private Endpoints, DNS privado y acceso desde red privada en ambos servicios.",
         "Arquitectura alternativa; no implementada en el Bicep actual."),
    ])
    paragraph(doc, "Para vaults nuevos se declara RBAC, protección contra purga, retención de 90 días "
              "y bypass AzureServices. Este bypass de servicios de confianza requiere justificar "
              "excepciones; no es un pase universal para Azure. disableLocalAuth=false es deliberado "
              "por la ruta de clave API. Regenerar la clave subyacente y actualizar su secreto "
              "en vault son acciones distintas. Una IP permitida no neutraliza una clave robada.")

    heading(doc, "6. Controles, líneas base y posición actual")
    paragraph(doc, "Referencias de diseño, no certificaciones: Microsoft Cloud Security Benchmark "
              "(dominios red, identidad, protección de datos, registro); Azure OpenAI security baseline "
              "(advierte posible desactualización de su versión); seguridad de red y registros de "
              "Key Vault; buenas prácticas de seguridad de IA Azure. Las recomendaciones se "
              "aplican solo en la medida documentada en código y plantillas.")
    table(doc, ("Dominio", "Demostrado por código/plantilla", "Brecha para el despliegue real"), [
        ("Red", "ACL/IP en recursos nuevos; dos perfiles públicos diferenciados.",
         "Verificar ambos firewalls, recursos reutilizados y egress del puesto."),
        ("Identidad", "Entra para vault; rol Secrets User opcional de solo lectura; escritores separados.",
         "Auditar RBAC efectivo, identidad del puesto y permisos del editor."),
        ("Datos", "Claves no se guardan en repositorio/argumentos; ejemplos sintéticos.",
         "Clasificar datos, residencia/SKU, retención y acceso de herramientas."),
        ("Protección", "Vault nuevo con purge protection + 90 días; validación de host ARM.",
         "Revisar excepción AzureServices y cifrado/estado del dispositivo."),
        ("Observabilidad", "Smoke-test pequeño y presupuesto opcional al 80 %.",
         "Diagnósticos/alertas no desplegados; probar recepción y revocación."),
        ("Gobierno", "What-if/aprobación; selección explícita de GlobalStandard.",
         "Aprobar propietario de riesgo, perfil de red y uso de datos."),
    ])
    paragraph(doc, "La conformidad con una baseline no se infiere de este cuadro. Los controles "
              "de contenido del servicio (raiPolicyName Microsoft.DefaultV2 en deployments "
              "declarados) tampoco sustituyen la clasificación de datos, la autorización de "
              "herramientas o una evaluación de riesgos de IA.")

    heading(doc, "7. Riesgos y pruebas de aceptación antes de producción")
    table(doc, ("Prioridad", "Riesgo / estado", "Evidencia mínima privada"), [
        ("Alta", "Proveedor/ruta de datos real no verificados (H07).",
         "CLI, subagentes, utilidades y VS Code por versión; TLS/destino; conflictos de proveedor."),
        ("Alta", "Red y observabilidad reales no verificadas (H08).",
         "Reglas de ambos servicios, RBAC, diagnósticos y alerta recibida por responsable."),
        ("Alta", "Fuga y revocación de clave; egress de herramientas.",
         "Canario ausente de logs; rotación de clave Azure y secreto vault; aislamiento de archivos."),
        ("Media", "Disponibilidad, costes y geografía de proceso.",
         "SKU autorizado, cuotas TPM/RPM y concurrencia; aceptación regional y presupuesto."),
    ])
    paragraph(doc, "Infraestructura: ejecutar P01 y P05 si selected_ips, o P06/P07 si any_ip; "
              "comprobar recursos reutilizados. Cliente: P02-P04, P08-P10, P12-P13 y P15; incluir "
              "T07, T12 y T16. Seguridad/operaciones: P14 y P16, respuesta y propietario con "
              "tiempo de revocación medible. P11 (renovación de token de inferencia) solo aplica "
              "a una futura migración a Entra para inferencia. Mocks no validan cloud ni "
              "persistencia del editor. No introducir datos de clientes hasta superar gates "
              "aplicables y aceptar explícitamente el riesgo residual del perfil escogido.")

    heading(doc, "8. Referencias y mantenimiento")
    for text in (
        "Código: infra/main.bicep; infra/resources.bicep; scripts/provision.py; "
        "scripts/foundry.py; scripts/seed_vault.py; scripts/vscode_config.py; tools/vscode-import/extension.js.",
        "Evaluación del repositorio y pruebas: docs/es/06-security.md; "
        "docs/es/02-create-azure-openai-deployment.md; docs/es/05-subagents.md.",
        "Guías externas locales aportadas: Guia_acceso_publico_Copilot_Azure_OpenAI.docx "
        "e Informe_seguridad_Copilot_Azure_OpenAI.docx. No forman parte de esta publicación.",
        "Microsoft Cloud Security Benchmark: https://learn.microsoft.com/security/benchmark/azure/",
        "Azure OpenAI security baseline: "
        "https://learn.microsoft.com/security/benchmark/azure/baselines/azure-openai-security-baseline",
        "Red de Azure AI: https://learn.microsoft.com/azure/ai-services/cognitive-services-virtual-networks",
        "Red de Key Vault: https://learn.microsoft.com/azure/key-vault/general/network-security",
        "Registros de Key Vault: https://learn.microsoft.com/azure/key-vault/general/logging",
        "Seguridad de IA Azure: https://learn.microsoft.com/azure/security/fundamentals/ai-security-best-practices",
    ):
        paragraph(doc, text, "List Bullet")
    paragraph(doc, "Actualizar este Word y la evaluación Markdown al modificar el cliente, Bicep, "
              "perfil de red o recomendaciones oficiales. Conservar evidencia operativa fuera del "
              "repositorio; no añadir capturas, nombres, identificadores o metadatos reales.")

    for relationships, suffix in ((doc.part.rels, "/customXml"),
                                  (doc.part.package.rels, "/thumbnail")):
        for rel in list(relationships.values()):
            if rel.reltype.endswith(suffix):
                relationships.pop(rel.rId)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Generated {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
